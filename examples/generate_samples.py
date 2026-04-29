#!/usr/bin/env python3
"""生成 AIDOC 示例文件：Word、PDF、图片"""
import os, sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))  # examples/
os.chdir(SCRIPT_DIR)

# ═══ Word 文档 ═══
print("📄 生成 Word...")
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.add_heading('投行底稿数字平台 — 项目需求说明书', 0)
doc.add_paragraph('版本：v1.0 | 日期：2026-04-25 | 编制：金融科技部', style='Subtitle')

doc.add_heading('一、项目背景', 1)
doc.add_paragraph('当前投行底稿管理存在以下痛点：')
for item in ['底稿文件分散存储，检索困难', '监管检查时调阅效率低', '版本管理混乱，无法追溯', '缺乏统一文件命名规范']:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('二、建设目标', 1)
doc.add_paragraph('打造覆盖底稿生成、流转、归档、检查全流程的数字平台，实现底稿管理的标准化、数字化和智能化。')

doc.add_heading('三、功能需求', 1)
doc.add_heading('3.1 底稿录入', 2)
doc.add_paragraph('支持 Word、PDF、扫描件等多种格式上传，自动提取元数据，支持批量导入。')
doc.add_heading('3.2 智能检索', 2)
doc.add_paragraph('基于全文检索的底稿搜索，支持按项目、类型、日期等多维度筛选，支持 OCR 识别扫描件文字。')
doc.add_heading('3.3 版本管理', 2)
doc.add_paragraph('每次修改自动生成新版本，支持版本对比和回溯，操作记录完整可追溯。')
doc.add_heading('3.4 监管调阅', 2)
doc.add_paragraph('一键打包导出底稿包，生成调阅清单，支持水印保护。')

doc.add_heading('四、技术方案', 1)
table = doc.add_table(rows=6, cols=3)
table.style = 'Light Shading'
cells = [
    ('模块', '技术选型', '说明'),
    ('前端', 'Vue 3 + Element Plus', '管理后台'),
    ('后端', 'Spring Boot 3', '微服务架构'),
    ('存储', 'MinIO + PostgreSQL', '文件及元数据'),
    ('搜索', 'Elasticsearch', '全文检索'),
    ('OCR', 'PaddleOCR', '扫描件文字识别'),
]
for i, row_data in enumerate(cells):
    for j, cell_text in enumerate(row_data):
        table.cell(i, j).text = cell_text

doc.add_heading('五、项目计划', 1)
plan_items = [
    ('2026.05', '需求评审'),
    ('2026.06', '系统设计'),
    ('2026.07-08', '开发阶段'),
    ('2026.09', '测试阶段'),
    ('2026.10', '正式上线'),
]
for time, task in plan_items:
    doc.add_paragraph(f'{time}  {task}')

doc.add_paragraph('审批人：张总')
doc.save('项目需求说明书.docx')
print('  ✅ 项目需求说明书.docx')

# ═══ PDF 文档 ═══
print("📄 生成 PDF...")
from fpdf import FPDF

pdf = FPDF()
pdf.add_page()
pdf.set_font('Helvetica', 'B', 18)
pdf.cell(0, 12, 'Investment Banking Document Platform', ln=True, align='C')
pdf.set_font('Helvetica', '', 10)
pdf.cell(0, 8, 'Requirements Specification v1.0 | 2026-04-25', ln=True, align='C')
pdf.ln(10)

pdf.set_font('Helvetica', 'B', 14)
pdf.cell(0, 10, '1. Background', ln=True)
pdf.set_font('Helvetica', '', 11)
pdf.multi_cell(0, 6, 'Current pain points in investment banking document management:\n'
    '- Documents stored across multiple locations, hard to find\n'
    '- Low efficiency when responding to regulatory inspections\n'
    '- No version control, hard to track changes\n'
    '- No unified naming conventions')

pdf.set_font('Helvetica', 'B', 14)
pdf.cell(0, 10, '2. Technical Architecture', ln=True)
pdf.set_font('Helvetica', '', 11)
pdf.multi_cell(0, 6, 'Frontend: Vue 3 + Element Plus\nBackend: Spring Boot 3\nStorage: MinIO + PostgreSQL\nSearch: Elasticsearch\nOCR: PaddleOCR')

pdf.set_font('Helvetica', 'B', 14)
pdf.cell(0, 10, '3. Timeline', ln=True)
pdf.set_font('Helvetica', '', 11)
pdf.multi_cell(0, 6, 'May 2026 - Requirements Review\nJun 2026 - System Design\nJul-Aug 2026 - Development\nSep 2026 - Testing\nOct 2026 - Launch')

pdf.output('项目需求说明书-英文版.pdf')
print('  ✅ 项目需求说明书-英文版.pdf')

# ═══ 图片 ═══
print("🖼️ 生成图片...")
from PIL import Image, ImageDraw, ImageFont

def create_chart(filename, title, labels, values, color):
    w, h = 600, 400
    img = Image.new('RGB', (w, h), 'white')
    draw = ImageDraw.Draw(img)
    
    # 标题
    try:
        font_title = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 22)
        font_label = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 14)
    except:
        font_title = ImageFont.load_default()
        font_label = ImageFont.load_default()
    
    draw.text((w//2, 20), title, fill='#1d1d1f', anchor='mt', font=font_title)
    
    # 柱状图
    bar_w = 60
    max_h = 200
    start_x = (w - len(labels) * (bar_w + 30)) // 2
    baseline = h - 60
    max_val = max(values) if values else 1
    
    for i, (label, val) in enumerate(zip(labels, values)):
        bar_h = int(val / max_val * max_h)
        x = start_x + i * (bar_w + 30)
        draw.rectangle([x, baseline - bar_h, x + bar_w, baseline],
                       fill=color, outline='#d2d2d7')
        draw.text((x + bar_w//2, baseline + 10), str(val),
                  fill='#1d1d1f', anchor='mt', font=font_label)
        # 多行标签
        lines = label.split('\n')
        for li, line in enumerate(lines):
            draw.text((x + bar_w//2, baseline + 30 + li * 18), line,
                      fill='#86868b', anchor='mt', font=font_label)
    
    img.save(filename)
    print(f'  ✅ {filename}')

create_chart('技术栈架构图.png',
    '投行底稿平台技术栈',
    ['Vue 3', 'Spring\nBoot 3', 'MinIO', 'PostgreSQL', 'Elastic-\nsearch', 'PaddleOCR'],
    [85, 90, 70, 75, 80, 65],
    '#007aff')

create_chart('项目里程碑.png',
    '项目里程碑计划',
    ['需求评审', '系统设计', '开发阶段', '测试阶段', '正式上线'],
    [30, 50, 85, 75, 100],
    '#34c759')

print()
print('🎉 全部生成完成！')